import json, re
from pathlib import Path
import openpyxl
root=Path(__file__).resolve().parent.parent

def sheet(name):
    return openpyxl.load_workbook(root/'sources'/name, read_only=True, data_only=True).active

def year(v):
    match=re.match(r'^(19\d\d|20\d\d)',str(v))
    return int(match.group(1)) if match else None

def num(v):
    return round(float(v),1) if isinstance(v,(int,float)) else None

s=sheet('insee-ratios-2026-08.xlsx'); ratios=[]
for c in range(2,68):
    y=year(s.cell(4,c).value)
    if y is None: continue
    r={'year':y,'revenue':num(s.cell(7,c).value),'spending':num(s.cell(5,c).value),'balance':num(s.cell(9,c).value),'debt':num(s.cell(10,c).value)}
    if y==2024:r['status']='semi-définitif'
    if y==2025:r['status']='provisoire'
    ratios.append(r)
(root/'sources/national-ratios-insee.json').write_text(json.dumps({'source':'https://www.insee.fr/fr/statistiques/2381414','publication':'2026-08-28','unit':'% du PIB','years':ratios},ensure_ascii=False,indent=2)+'\n',encoding='utf8')

s=sheet('insee-prelevements-2026-08.xlsx'); levies={}
for i in range(5,72):
    y=year(s.cell(i,1).value); v=num(s.cell(i,2).value)
    if y and v is not None: levies[y]=v
s=sheet('insee-impots-2026-08.xlsx')
categories=[(5,'vat','TVA'),(6,'imports','Impôts et droits sur les importations'),(7,'products','Autres impôts sur les produits'),(11,'payroll','Impôts sur les salaires et la main-d’œuvre'),(13,'production','Impôts divers sur la production'),(20,'income','Impôts courants sur le revenu'),(26,'otherCurrent','Autres impôts courants'),(30,'capital','Impôts en capital')]
details=[(22,'income','dont CSG hors jeux'),(23,'income','dont CRDS hors jeux'),(24,'income','dont impôt sur le revenu, avant crédits d’impôt'),(25,'income','dont impôt sur les sociétés, avant crédits d’impôt'),(15,'production','dont taxe foncière'),(9,'products','dont accises sur les produits énergétiques (TICPE)')]
columns={year(s.cell(4,c).value):c for c in range(2,33)}; years=[]
for y in sorted(levies):
    c=columns.get(y); cats=[]; detail=[]
    if c:
        cats=[{'id':key,'label':label,'value':num(s.cell(r,c).value)} for r,key,label in categories]
        detail=[{'parent':parent,'label':label,'value':num(s.cell(r,c).value)} for r,parent,label in details]
        if any(v['value'] is None for v in cats+detail): raise ValueError(f'Missing tax value: {y}')
    years.append({'year':y,'levyRate':levies[y],'categories':cats,'details':detail})
(root/'sources/national-taxes-insee.json').write_text(json.dumps({'publication':'2026-08-28','source':'https://www.insee.fr/fr/statistiques/2381408','levySource':'https://www.insee.fr/fr/statistiques/2381412','unit':'milliards euros courants','years':years},ensure_ascii=False,indent=2)+'\n',encoding='utf8')
