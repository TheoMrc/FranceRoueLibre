"""Extract annual national account totals and non-overlapping subtotals from Insee table 3.201."""
import json
from pathlib import Path
import openpyxl

root = Path(__file__).resolve().parent.parent
book = openpyxl.load_workbook(root / 'sources/insee-apu-3201-2026-08.xlsx', data_only=True, read_only=True)
sheet = book['t_3201']
columns = {sheet.cell(3, col).value: col for col in range(3, sheet.max_column + 1)}
spending_rows = [(5, 'Fonctionnement'), (12, 'Intérêts'), (15, 'Prestations et autres transferts'), (23, 'Acquisitions nettes d’actifs non financiers')]
revenue_rows = [(34, 'Recettes de production'), (40, 'Revenus de la propriété'), (45, 'Impôts et cotisations sociales'), (53, 'Autres transferts')]

def value(row, col):
    result = sheet.cell(row, col).value
    if not isinstance(result, (int, float)):
        raise ValueError(f'Missing numeric value at {row}, {col}')
    return round(result, 4)

years = []
for year in range(1960, 2026):
    col = columns[year]
    expenses = value(27, col)
    receipts = value(58, col)
    balance = value(66, col)
    spending = [{'label': label, 'value': value(row, col)} for row, label in spending_rows]
    revenue = [{'label': label, 'value': value(row, col)} for row, label in revenue_rows]
    if abs(sum(item['value'] for item in spending) - expenses) > .001 or abs(sum(item['value'] for item in revenue) - receipts) > .001:
        raise ValueError(f'Breakdown mismatch in {year}')
    if abs(receipts - expenses - balance) > .001:
        raise ValueError(f'Balance mismatch in {year}')
    years.append({'year': year, 'spending': expenses, 'revenue': receipts, 'balance': balance, 'spendingBreakdown': spending, 'revenueBreakdown': revenue})
out = {'source': 'https://www.insee.fr/fr/statistiques/8988845', 'table': '3.201', 'publication': '2026-08-28', 'unit': 'milliards euros courants', 'years': years}
(root / 'sources/national-accounts-3201.json').write_text(json.dumps(out, ensure_ascii=False, indent=2) + '\n', encoding='utf8')
