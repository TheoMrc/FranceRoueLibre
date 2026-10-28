"""Extract Maastricht debt and deficit-to-debt bridge from Insee tables 3.101 and 3.105."""
import json
from pathlib import Path
import openpyxl

root = Path(__file__).resolve().parent.parent
stocks = openpyxl.load_workbook(root / 'sources/insee-dette-3101-2026-08.xlsx', data_only=True, read_only=True)['T_3101']
bridge = openpyxl.load_workbook(root / 'sources/insee-dette-3105-2026-08.xlsx', data_only=True, read_only=True)['t_3105']
stock_cols = {str(stocks.cell(4, c).value): c for c in range(3, stocks.max_column + 1)}
bridge_cols = {str(bridge.cell(3, c).value): c for c in range(3, bridge.max_column + 1)}
sectors = [(9, 'État'), (10, 'Organismes divers d’administration centrale'), (11, 'Administrations publiques locales'), (12, 'Administrations de sécurité sociale')]
adjustments = [(7, 'Acquisitions nettes d’actifs financiers'), (15, 'Variation des passifs hors dette de Maastricht'), (19, 'Changements de volume de la dette'), (24, 'Solde des comptes à recevoir et à payer'), (26, 'Écart statistique entre comptes financiers et non financiers')]

def number(sheet, row, col):
    val = sheet.cell(row, col).value
    if not isinstance(val, (int, float)):
        raise ValueError(f'Non-numeric value at {sheet.title}:{row},{col}')
    return round(val, 4)

years = []
for year in range(1978, 2026):
    col = stock_cols[str(year)]
    debt = number(stocks, 7, col)
    previous = number(stocks, 7, stock_cols[str(year - 1)]) if str(year - 1) in stock_cols else None
    values = [{'label': label, 'value': number(stocks, r, col)} for r, label in sectors]
    if abs(sum(x['value'] for x in values) - debt) > .21:
        raise ValueError(f'Sector contributions mismatch for {year}')
    row = {'year': year, 'debt': debt, 'previousDebt': previous, 'sectors': values, 'deficit': None, 'adjustments': [], 'change': None}
    if str(year) in bridge_cols:
        bcol = bridge_cols[str(year)]
        components = [{'label': label, 'value': number(bridge, r, bcol)} for r, label in adjustments]
        deficit = number(bridge, 5, bcol); change = number(bridge, 28, bcol)
        if abs(deficit + sum(x['value'] for x in components) - change) > .11:
            raise ValueError(f'Bridge mismatch for {year}')
        # Preserve cross-table differences for the interface to report.
        row['stockDifference'] = round(debt - previous - change, 4) if previous is not None else None
        row.update(deficit=deficit, adjustments=components, change=change)
    years.append(row)
output = {'source': 'https://www.insee.fr/fr/statistiques/8988843', 'tables': ['3.101', '3.105'], 'publication': '2026-08-28', 'unit': 'milliards euros courants', 'years': years}
(root / 'sources/national-debt.json').write_text(json.dumps(output, ensure_ascii=False, indent=2) + '\n', encoding='utf8')
