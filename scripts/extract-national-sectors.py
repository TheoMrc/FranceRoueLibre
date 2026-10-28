"""Extract Insee 3.215 sector totals, preserving published consolidation differences."""
import json
from pathlib import Path
import openpyxl

root = Path(__file__).resolve().parent.parent
sheet = openpyxl.load_workbook(root / 'sources/insee-apu-3215-2026-08.xlsx', data_only=True, read_only=True)['t_3215_fr']
columns = {sheet.cell(3, col).value: col for col in range(3, sheet.max_column + 1)}
sectors = [('state', 'État', 7, 8), ('odac', 'Organismes divers d’administration centrale', 10, 11), ('local', 'Administrations publiques locales', 13, 14), ('social', 'Administrations de sécurité sociale', 16, 17)]

def number(row, col):
    val = sheet.cell(row, col).value
    if not isinstance(val, (int, float)):
        raise ValueError(f'Missing number at row {row}, column {col}')
    return round(val, 4)

years = []
for year in range(1978, 2026):
    col = columns[year]
    rows = []
    for key, label, spending_row, revenue_row in sectors:
        spending, revenue = number(spending_row, col), number(revenue_row, col)
        rows.append({'id': key, 'label': label, 'spending': spending, 'revenue': revenue, 'balance': round(revenue - spending, 4)})
    years.append({'year': year, 'sectors': rows, 'totalSpending': number(20, col), 'totalRevenue': number(21, col)})
output = {'source': 'https://www.insee.fr/fr/statistiques/8988845', 'table': '3.215', 'publication': '2026-08-28', 'unit': 'milliards euros courants', 'years': years}
(root / 'sources/national-sectors-3215.json').write_text(json.dumps(output, ensure_ascii=False, indent=2) + '\n', encoding='utf8')
